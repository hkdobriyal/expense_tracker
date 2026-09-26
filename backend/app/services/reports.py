"""Report builders. Each returns a title, column headers and rows of plain values
so the same data can be exported as CSV, XLSX or rendered for printing (PDF via the browser)."""

from __future__ import annotations

import csv
import io
from datetime import date

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..models import Account, Transaction
from ..money import from_minor
from . import budgets as budget_svc
from . import ledger, networth
from .context import UserContext

REPORT_TYPES = ("monthly_summary", "transactions", "expenses", "income", "budgets", "cash_flow", "net_worth", "savings")


def _m(minor: int | None, currency: str):
    # Decimal, not float: exports keep exact paise (openpyxl and csv both handle Decimal).
    return from_minor(minor, currency) if minor is not None else None


def build(db: Session, ctx: UserContext, report: str, start: date, end: date) -> dict:
    cur = ctx.base_currency
    if report == "transactions":
        txns = db.scalars(
            select(Transaction).where(Transaction.user_id == ctx.user_id, Transaction.date >= start, Transaction.date <= end)
            .order_by(Transaction.date, Transaction.id).options(selectinload(Transaction.tags))
        ).all()
        return {
            "title": "Transactions", "columns": ["Date", "Type", "Account", "Description", "Merchant", "Category", "Amount", "Currency", f"Amount ({cur})", "Payment method", "Tags", "Notes", "Source", "Reviewed"],
            "rows": [[t.date.isoformat(), t.type, t.account.name, t.description, t.merchant.name if t.merchant else "", t.category.name if t.category else "",
                      _m(t.amount_minor, t.currency), t.currency, _m(t.base_amount_minor, cur), t.payment_method, ", ".join(x.name for x in t.tags), t.notes, t.source, t.reviewed] for t in txns],
        }
    if report in ("expenses", "income"):
        rows = []
        for group in ledger.category_breakdown(db, ctx.user_id, start, end, "expense" if report == "expenses" else "income"):
            rows.append([group["name"], "", _m(group["amount"], cur), group["share"]])
            for child in group["children"]:
                rows.append([group["name"], child["name"], _m(child["amount"], cur), None])
        return {"title": "Expense report" if report == "expenses" else "Income report", "columns": ["Category", "Subcategory", f"Amount ({cur})", "Share %"], "rows": rows}
    if report == "budgets":
        rows = [[b.name, b.category.name if b.category else "All spending", b.period, st["window_start"], st["window_end"], _m(st["limit"], cur), _m(st["spent"], cur), _m(st["remaining"], cur), st["usage_pct"], st["status"]]
                for b, st in budget_svc.all_budget_statuses(db, ctx.user_id, end, ctx.week_start)]
        return {"title": "Budget report", "columns": ["Budget", "Category", "Period", "From", "To", "Limit", "Spent", "Remaining", "Used %", "Status"], "rows": rows}
    if report == "cash_flow":
        cf = ledger.cash_flow(db, ctx.user_id, cur, start, end)
        series = ledger.bucketed_series(db, ctx.user_id, start, end, "month")
        rows = [[p["bucket"][:7], _m(p["income"], cur), _m(p["expenses"], cur), _m(p["invested"], cur), _m(p["net"], cur)] for p in series]
        rows.append(["Total", _m(cf["income"], cur), _m(cf["expenses"], cur), _m(cf["invested"], cur), _m(cf["net_cash_flow"], cur)])
        return {"title": "Cash flow report", "columns": ["Month", "Income", "Expenses", "Invested", "Net cash flow"], "rows": rows,
                "summary": {"Opening balance": _m(cf["opening_balance"], cur), "Closing balance": _m(cf["closing_balance"], cur)}}
    if report == "net_worth":
        months = max(1, (end.year - start.year) * 12 + end.month - start.month + 1)
        hist = networth.history(db, ctx.user_id, cur, end, min(months, 60))
        rows = [[p["date"], _m(p["assets"], cur), _m(p["liabilities"], cur), _m(p["net_worth"], cur)] for p in hist]
        accounts = db.scalars(select(Account).where(Account.user_id == ctx.user_id)).all()
        by_id = {a.id: a for a in accounts}
        breakdown = [[by_id[b.account_id].name, by_id[b.account_id].type, _m(b.balance_minor, b.currency), b.currency, _m(b.base_minor, cur)] for b in ledger.balances_in_base(db, ctx.user_id, cur, end, accounts)]
        return {"title": "Net worth report", "columns": ["Date", "Assets", "Liabilities", "Net worth"], "rows": rows,
                "extra": {"title": "Accounts", "columns": ["Account", "Type", "Balance", "Currency", f"Balance ({cur})"], "rows": breakdown}}
    if report == "savings":
        series = ledger.bucketed_series(db, ctx.user_id, start, end, "month")
        rows = [[p["bucket"][:7], _m(p["income"], cur), _m(p["expenses"], cur), _m(p["savings"], cur), round(p["savings"] / p["income"] * 100, 1) if p["income"] > 0 else None] for p in series]
        return {"title": "Savings report", "columns": ["Month", "Income", "Expenses", "Saved", "Savings rate %"], "rows": rows}
    # monthly_summary
    totals = ledger.period_totals(db, ctx.user_id, start, end)
    rows = [["Income", _m(totals.income, cur)], ["Expenses", _m(totals.expenses, cur)], ["Refunds (included above)", _m(totals.refunds, cur)],
            ["Invested", _m(totals.invested, cur)], ["Saved", _m(totals.savings, cur)], ["Savings rate %", totals.savings_rate],
            ["Net cash flow", _m(totals.net_cash_flow, cur)], ["Transactions", totals.transaction_count]]
    top = ledger.category_breakdown(db, ctx.user_id, start, end)[:10]
    return {"title": "Monthly summary", "columns": ["Metric", "Value"], "rows": rows,
            "extra": {"title": "Top categories", "columns": ["Category", f"Amount ({cur})", "Share %"], "rows": [[g["name"], _m(g["amount"], cur), g["share"]] for g in top]}}


def to_csv(report: dict) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(report["columns"])
    writer.writerows(report["rows"])
    if report.get("extra"):
        writer.writerow([])
        writer.writerow([report["extra"]["title"]])
        writer.writerow(report["extra"]["columns"])
        writer.writerows(report["extra"]["rows"])
    return buf.getvalue()


def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31]
    ws.append(report["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in report["rows"]:
        ws.append(row)
    if report.get("extra"):
        extra = wb.create_sheet(report["extra"]["title"][:31])
        extra.append(report["extra"]["columns"])
        for cell in extra[1]:
            cell.font = Font(bold=True)
        for row in report["extra"]["rows"]:
            extra.append(row)
    for sheet in wb.worksheets:
        for column in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
