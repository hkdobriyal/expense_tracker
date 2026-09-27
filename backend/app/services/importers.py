"""Statement file parsers: CSV/TSV/TXT, XLS/XLSX (and HTML-as-XLS), OFX/QFX, JSON, DOCX,
PDF (tables, text or scanned via OCR) and images (OCR).

Every parser produces the same shape – ``headers`` + ``rows`` of strings – so
one column-mapping step handles all formats. Mapping turns rows into
normalised transactions (date, description, signed amount, …).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from ..money import MoneyError, to_decimal

MAX_ROWS = 5000

DATE_FORMATS = (
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y",
    "%d %b %y", "%d/%b/%Y", "%b %d, %Y", "%m/%d/%Y", "%Y/%m/%d", "%Y%m%d", "%d %B %Y",
)

BANK_MARKERS = [
    ("hdfc", "HDFC Bank"), ("state bank of india", "State Bank of India"), ("sbi", "State Bank of India"), ("icici", "ICICI Bank"),
    ("axis", "Axis Bank"), ("kotak", "Kotak Mahindra Bank"), ("yes bank", "Yes Bank"), ("indusind", "IndusInd Bank"),
    ("punjab national", "Punjab National Bank"), ("bank of baroda", "Bank of Baroda"), ("canara", "Canara Bank"),
    ("idfc", "IDFC FIRST Bank"), ("federal bank", "Federal Bank"), ("au small finance", "AU Small Finance Bank"),
]


class ImportParseError(ValueError):
    pass


@dataclass
class ParsedFile:
    file_format: str
    headers: list[str]
    rows: list[list[str]]
    bank_detected: str = ""
    preset_mapping: dict | None = None  # formats with a fixed layout (OFX/PDF) need no user mapping
    method: str = ""  # how rows were obtained: csv, pdf-table, pdf-text, pdf-ocr, image-ocr, docx-table…
    notes: list | None = None


SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xls", ".xlsx", ".xlsm", ".ofx", ".qfx", ".pdf", ".docx", ".json", ".html", ".htm", ".png", ".jpg", ".jpeg", ".webp")


def detect_format(filename: str, content: bytes) -> str:
    """Decide by content first (magic bytes), then by extension."""
    import zipfile

    name = filename.lower()
    head = content[:1024].lstrip()
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if head[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "webp"
    if name.endswith((".ofx", ".qfx")) or b"<OFX>" in content[:4096].upper() or head.startswith(b"OFXHEADER"):
        return "qfx" if name.endswith(".qfx") else "ofx"
    if head.startswith(b"PK\x03\x04"):
        try:
            names = zipfile.ZipFile(io.BytesIO(content)).namelist()
        except zipfile.BadZipFile as exc:
            raise ImportParseError("The file looks like a damaged Office document.") from exc
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        raise ImportParseError("Unsupported Office/zip file. Upload XLSX, DOCX, CSV, PDF or OFX.")
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or name.endswith(".xls"):
        return "xls"
    if name.endswith(".doc"):
        raise ImportParseError("Old .doc files aren't supported - save it as .docx or PDF first.")
    if name.endswith(".json") or head[:1] in (b"[", b"{"):
        return "json"
    if name.endswith((".html", ".htm")) or head.lower().startswith((b"<html", b"<!doctype html", b"<table")):
        return "html"
    if name.endswith((".csv", ".tsv")):
        return "csv"
    if name.endswith(".txt") or "." not in name or _looks_like_text(content):
        return "txt"
    raise ImportParseError("Unsupported file type. Supported: CSV, TSV, TXT, XLS, XLSX, OFX/QFX, PDF, DOCX, JSON, HTML, PNG/JPG.")


def _looks_like_text(content: bytes) -> bool:
    sample = content[:2000]
    return bool(sample) and all(b >= 32 or b in (9, 10, 13) for b in sample)


def detect_bank(text: str) -> str:
    lower = text.lower()
    return next((label for marker, label in BANK_MARKERS if marker in lower), "")


def _is_header_row(cells: list[str]) -> bool:
    joined = " ".join(c.lower() for c in cells)
    has_date = "date" in joined
    has_text = any(k in joined for k in ("narration", "particular", "description", "remark", "details", "transaction"))
    has_amount = any(k in joined for k in ("amount", "debit", "credit", "withdrawal", "deposit", "dr", "cr"))
    return has_date and (has_text or has_amount)


def _tabulate(rows: list[list[str]], file_format: str, bank: str) -> ParsedFile:
    rows = [[(c if c is not None else "").strip() for c in r] for r in rows]
    header_idx = next((i for i, r in enumerate(rows[:40]) if _is_header_row(r)), None)
    if header_idx is None:
        raise ImportParseError("Could not find a header row with a date and an amount/description column.")
    headers = [h or f"Column {i + 1}" for i, h in enumerate(rows[header_idx])]
    width = len(headers)
    body = []
    for r in rows[header_idx + 1:]:
        if not any(r):
            continue
        body.append((r + [""] * width)[:width])
        if len(body) > MAX_ROWS:
            raise ImportParseError(f"Files are limited to {MAX_ROWS} rows. Split the statement and import it in parts.")
    return ParsedFile(file_format, headers, body, bank)


def parse_csv(content: bytes) -> ParsedFile:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _tabulate(rows, "csv", detect_bank(sample))


def parse_xlsx(content: bytes) -> ParsedFile:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # corrupt / not an xlsx
        raise ImportParseError(f"Could not open the spreadsheet: {exc}") from exc
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = []
        for v in row:
            if isinstance(v, datetime):
                cells.append(v.date().isoformat())
            elif isinstance(v, date):
                cells.append(v.isoformat())
            elif isinstance(v, float):
                cells.append(format(Decimal(repr(v)), "f"))
            else:
                cells.append("" if v is None else str(v))
        rows.append(cells)
        if len(rows) > MAX_ROWS + 50:
            break
    wb.close()
    return _tabulate(rows, "xlsx", detect_bank(" ".join(" ".join(r) for r in rows[:15])))


_OFX_TXN = re.compile(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|(?=</BANKTRANLIST>))", re.S | re.I)


def _ofx_field(block: str, tag: str) -> str:
    m = re.search(rf"<{tag}>([^<\r\n]*)", block, re.I)
    return m.group(1).strip() if m else ""


def parse_ofx(content: bytes, file_format: str = "ofx") -> ParsedFile:
    text = content.decode("utf-8", errors="replace")
    rows = []
    for block in _OFX_TXN.findall(text):
        posted = _ofx_field(block, "DTPOSTED")[:8]
        rows.append([
            f"{posted[:4]}-{posted[4:6]}-{posted[6:8]}" if len(posted) == 8 else posted,
            _ofx_field(block, "NAME") or _ofx_field(block, "MEMO"),
            _ofx_field(block, "TRNAMT"),
            _ofx_field(block, "FITID"),
            _ofx_field(block, "MEMO"),
            _ofx_field(block, "TRNTYPE"),
        ])
    if not rows:
        raise ImportParseError("No transactions (<STMTTRN>) found in the OFX file.")
    org = _ofx_field(text, "ORG")
    headers = ["Date", "Description", "Amount", "FITID", "Memo", "Type"]
    mapping = {"date": 0, "description": 1, "amount": 2, "external_id": 3, "notes": 4, "date_format": "%Y-%m-%d", "sign": "signed"}
    return ParsedFile(file_format, headers, rows[:MAX_ROWS], org or detect_bank(text[:3000]), mapping)


TEXT_HEADERS = ["Date", "Narration", "Amount", "Dr/Cr", "Balance"]
TEXT_MAPPING = {"date": 0, "description": 1, "amount": 2, "direction": 3, "balance": 4, "sign": "direction_column"}


def _from_text_lines(lines: list[str], file_format: str, bank_text: str, method: str) -> ParsedFile:
    from .statement_text import parse_lines

    rows, notes = parse_lines(lines)
    if not rows:
        raise ImportParseError(
            "No transaction lines were recognised. Each transaction should start with a date and include an amount. "
            "If your bank offers a CSV/Excel download, that imports most reliably."
        )
    parsed = ParsedFile(file_format, TEXT_HEADERS, rows[:MAX_ROWS], detect_bank(bank_text), dict(TEXT_MAPPING))
    parsed.method = method
    parsed.notes = notes
    return parsed


def _tables_to_parsed(tables: list[list[list[str]]], file_format: str, bank_text: str, method: str) -> ParsedFile | None:
    """Stitch tables (possibly split across pages) under the first header row found."""
    header_row: list[str] | None = None
    body: list[list[str]] = []
    for table in tables:
        cleaned = [[re.sub(r"\s+", " ", c or "").strip() for c in row] for row in table if row and any(c for c in row)]
        for row in cleaned:
            if header_row is None:
                if _is_header_row(row):
                    header_row = row
                continue
            if row == header_row or _is_header_row(row):
                continue  # header repeated on later pages
            body.append(row)
    if header_row is None or not body:
        return None
    parsed = _tabulate([header_row] + body, file_format, detect_bank(bank_text))
    parsed.method = method
    return parsed


def parse_pdf_any(content: bytes, password: str | None = None) -> ParsedFile:
    """PDF: tables → text lines → OCR, in that order."""
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        pdf = pdfplumber.open(io.BytesIO(content), password=password or "")
    except PDFPasswordIncorrect as exc:
        raise ImportParseError("This PDF is password protected. Enter the statement password (often your PAN in capitals, or name + date of birth)."
                               if not password else "Incorrect statement password.") from exc
    except Exception as exc:  # noqa: BLE001 - malformed PDF
        raise ImportParseError(f"Could not read the PDF: {exc}") from exc
    with pdf:
        pages = pdf.pages[:60]
        text = "\n".join((p.extract_text() or "") for p in pages)
        tables = [t for p in pages for t in (p.extract_tables() or [])]
        parsed = _tables_to_parsed(tables, "pdf", text[:4000], "pdf-table") if tables else None
        if parsed is not None and len(parsed.rows) >= 1:
            return parsed
        if text.strip():
            try:
                return _from_text_lines(text.splitlines(), "pdf", text[:4000], "pdf-text")
            except ImportParseError:
                pass
        from ..config import get_settings

        if not get_settings().ocr_enabled:
            raise ImportParseError("This PDF has no readable text (scanned image) and OCR is disabled (OCR_ENABLED=false).")
        lines: list[str] = []
        from .statement_text import ocr_image_lines

        for page in pages[:15]:
            lines += ocr_image_lines(page.to_image(resolution=200).original)
        return _from_text_lines(lines, "pdf", "\n".join(lines[:60]), "pdf-ocr")


def parse_image(content: bytes, file_format: str) -> ParsedFile:
    from PIL import Image, ImageOps

    from ..config import get_settings
    from .statement_text import ocr_image_lines

    if not get_settings().ocr_enabled:
        raise ImportParseError("Images need OCR, which is disabled (OCR_ENABLED=false).")
    try:
        image = ImageOps.exif_transpose(Image.open(io.BytesIO(content)))
    except Exception as exc:  # noqa: BLE001
        raise ImportParseError(f"Could not open the image: {exc}") from exc
    lines = ocr_image_lines(image)
    return _from_text_lines(lines, file_format, "\n".join(lines[:60]), "image-ocr")


def parse_docx(content: bytes) -> ParsedFile:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ImportParseError(f"Could not open the Word document: {exc}") from exc
    paragraphs = [p.text for p in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    bank_text = "\n".join(paragraphs[:40])
    parsed = _tables_to_parsed(tables, "docx", bank_text, "docx-table") if tables else None
    if parsed is not None:
        return parsed
    lines = paragraphs + [" ".join(row) for t in tables for row in t]
    return _from_text_lines(lines, "docx", bank_text, "docx-text")


class _HTMLTables:
    """Tiny HTML table reader (bank 'xls' downloads are often HTML)."""

    def __init__(self, text: str):
        from html.parser import HTMLParser

        self.tables: list[list[list[str]]] = []
        outer = self

        class P(HTMLParser):
            def __init__(self):
                super().__init__()
                self.row: list[str] | None = None
                self.cell: list[str] | None = None

            def handle_starttag(self, tag, attrs):
                if tag == "table":
                    outer.tables.append([])
                elif tag == "tr" and outer.tables:
                    self.row = []
                elif tag in ("td", "th") and self.row is not None:
                    self.cell = []

            def handle_endtag(self, tag):
                if tag in ("td", "th") and self.cell is not None and self.row is not None:
                    self.row.append(" ".join("".join(self.cell).split()))
                    self.cell = None
                elif tag == "tr" and self.row is not None and outer.tables:
                    outer.tables[-1].append(self.row)
                    self.row = None

            def handle_data(self, data):
                if self.cell is not None:
                    self.cell.append(data)

        P().feed(text)


def parse_html(content: bytes, file_format: str = "html") -> ParsedFile:
    text = content.decode("utf-8", errors="replace")
    tables = _HTMLTables(text).tables
    parsed = _tables_to_parsed(tables, file_format, re.sub(r"<[^>]+>", " ", text[:6000]), "html-table")
    if parsed is None:
        raise ImportParseError("No transaction table with a Date column was found in this file.")
    return parsed


def parse_xls(content: bytes) -> ParsedFile:
    if content[:8] != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        # Many banks' ".xls" downloads are really HTML or tab-separated text.
        head = content[:2000].lower()
        if b"<table" in head or b"<html" in head or b"<?xml" in head:
            return parse_html(content, "xls")
        return parse_csv(content)
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:  # noqa: BLE001
        raise ImportParseError(f"Could not open the .xls file: {exc}") from exc
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(min(sheet.nrows, MAX_ROWS + 60)):
        cells = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                cells.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                cells.append(format(Decimal(repr(cell.value)), "f"))
            else:
                cells.append(str(cell.value))
        rows.append(cells)
    return _tabulate(rows, "xls", detect_bank(" ".join(" ".join(r) for r in rows[:15])))


def parse_json(content: bytes) -> ParsedFile:
    import json

    try:
        data = json.loads(content.decode("utf-8-sig"))
    except ValueError as exc:
        raise ImportParseError(f"Invalid JSON: {exc}") from exc
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [])
    records = [r for r in data if isinstance(r, dict)] if isinstance(data, list) else []
    if not records:
        raise ImportParseError("The JSON file should contain a list of transaction objects.")
    headers: list[str] = []
    for r in records[:200]:
        headers += [k for k in r if k not in headers]
    rows = [["" if r.get(h) is None else str(r.get(h)) for h in headers] for r in records[:MAX_ROWS]]
    return _tabulate([headers] + rows, "json", "")


def parse_text_file(content: bytes) -> ParsedFile:
    """TXT: delimited table if it looks like one, otherwise free-text statement lines."""
    try:
        parsed = parse_csv(content)
        if len(parsed.headers) >= 3:  # a real delimited table
            parsed.file_format, parsed.method = "txt", "txt-table"
            return parsed
    except ImportParseError:
        pass
    text = content.decode("utf-8", errors="replace")
    return _from_text_lines(text.splitlines(), "txt", text[:3000], "text")


def parse_file(filename: str, content: bytes, password: str | None = None) -> ParsedFile:
    fmt = detect_format(filename, content)
    parsers = {
        "csv": lambda: parse_csv(content), "txt": lambda: parse_text_file(content), "xlsx": lambda: parse_xlsx(content),
        "xls": lambda: parse_xls(content), "ofx": lambda: parse_ofx(content, "ofx"), "qfx": lambda: parse_ofx(content, "qfx"),
        "pdf": lambda: parse_pdf_any(content, password), "docx": lambda: parse_docx(content), "html": lambda: parse_html(content),
        "json": lambda: parse_json(content), "png": lambda: parse_image(content, "png"), "jpg": lambda: parse_image(content, "jpg"),
        "webp": lambda: parse_image(content, "webp"),
    }
    parsed = parsers[fmt]()
    if not getattr(parsed, "method", None):
        parsed.method = fmt
    return parsed


# ---------------------------------------------------------------------------
# Column mapping
# ---------------------------------------------------------------------------

def suggest_mapping(headers: list[str]) -> dict:
    lower = [h.lower() for h in headers]

    def find(*keys, exclude=()):
        for i, h in enumerate(lower):
            if any(k in h for k in keys) and not any(x in h for x in exclude):
                return i
        return None

    mapping = {
        "date": find("txn date", "transaction date", "value date", "date", exclude=("chq",)),
        "description": find("narration", "particular", "description", "remark", "details"),
        "debit": find("withdrawal", "debit", exclude=("credit",)),
        "credit": find("deposit", "credit", exclude=("debit",)),
        "amount": find("amount", "amt", exclude=("balance",)),
        "direction": find("dr/cr", "cr/dr", "type"),
        "balance": find("balance"),
        "external_id": find("ref no", "reference", "chq", "utr", "transaction id"),
    }
    if mapping["debit"] is not None and mapping["credit"] is not None:
        mapping["sign"] = "debit_credit"
    elif mapping["direction"] is not None:
        mapping["sign"] = "direction_column"
    else:
        mapping["sign"] = "signed"
    mapping["date_format"] = None  # auto-detect
    return mapping


def detect_date_format(values: list[str]) -> str | None:
    samples = [v.strip().split(" ")[0] if re.match(r"^\d", v.strip()) and ":" in v else v.strip() for v in values if v and v.strip()][:200]
    if not samples:
        return None
    best, best_hits = None, 0
    for fmt in DATE_FORMATS:  # day-first formats come first: Indian statements are DD/MM
        hits = 0
        for s in samples:
            try:
                datetime.strptime(s, fmt)
                hits += 1
            except ValueError:
                pass
        if hits > best_hits:
            best, best_hits = fmt, hits
        if hits == len(samples):
            return fmt
    return best


def parse_date(value: str, fmt: str | None) -> date:
    raw = value.strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}[T ]", raw):
        raw = raw[:10]
    elif ":" in raw and " " in raw:
        raw = raw.split(" ")[0]
    formats = [fmt] if fmt else list(DATE_FORMATS)
    for f in formats:
        try:
            return datetime.strptime(raw, f).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{value}'")


def parse_amount(value: str) -> Decimal | None:
    raw = (value or "").strip()
    if not raw or raw in ("-", "--"):
        return None
    negative = raw.startswith("(") and raw.endswith(")") or raw.startswith("-") or raw.endswith("-")
    cleaned = re.sub(r"(?i)(inr|rs\.?|₹|cr|dr|[()\s+-])", "", raw)
    if not cleaned:
        return None
    try:
        amount = to_decimal(cleaned)
    except MoneyError:
        raise ValueError(f"Unrecognised amount '{value}'")
    return -amount if negative else amount


def guess_merchant(narration: str) -> str | None:
    from .extraction import extract

    return extract(narration).get("merchant")


def guess_payment_method(narration: str) -> str:
    from .extraction import extract, payment_method

    return payment_method(extract(narration))


def guess_type(narration: str, direction: str) -> str:
    upper = narration.upper()
    if direction == "in":
        if "REFUND" in upper or "REVERSAL" in upper or "RETURN" in upper:
            return "refund"
        return "income"
    if any(k in upper for k in ("SIP", "MUTUAL FUND", "ZERODHA", "GROWW", "KFINTECH", "CAMS", "INDMONEY", "UPSTOX")):
        return "investment"
    return "expense"


def normalize_rows(headers: list[str], rows: list[list[str]], mapping: dict) -> list[dict]:
    """Apply a column mapping. Each output row has either the fields or an ``error``."""

    def col(row: list[str], key: str) -> str:
        idx = mapping.get(key)
        if idx is None or idx == "" or int(idx) >= len(row):
            return ""
        return row[int(idx)]

    from .extraction import extract, payment_method

    date_fmt = mapping.get("date_format") or detect_date_format([col(r, "date") for r in rows])
    sign_mode = mapping.get("sign", "signed")
    invert = bool(mapping.get("invert"))
    out = []
    for index, row in enumerate(rows):
        item: dict = {"row": index, "raw": row}
        try:
            if mapping.get("date") is None:
                raise ValueError("Map the date column")
            try:
                day = parse_date(col(row, "date"), date_fmt)
            except ValueError:
                if mapping.get("date_format"):
                    raise  # the user chose a format explicitly
                day = parse_date(col(row, "date"), None)  # statements printed as text often mix date styles
            description = re.sub(r"\s+", " ", col(row, "description")).strip()
            if not description:
                raise ValueError("Empty description")
            if sign_mode == "debit_credit":
                debit, credit = parse_amount(col(row, "debit")), parse_amount(col(row, "credit"))
                if credit and credit != 0:
                    signed = abs(credit)
                elif debit and debit != 0:
                    signed = -abs(debit)
                else:
                    raise ValueError("No debit or credit amount")
            else:
                amount = parse_amount(col(row, "amount"))
                if amount is None or amount == 0:
                    raise ValueError("No amount")
                if sign_mode == "direction_column":
                    d = col(row, "direction").strip().upper()
                    is_credit = d.startswith("CR") or d.startswith("C") or d in ("DEPOSIT", "CREDIT", "IN")
                    signed = abs(amount) if is_credit else -abs(amount)
                else:
                    signed = amount
            if invert:
                signed = -signed
            direction = "in" if signed > 0 else "out"
            entities = extract(description)
            item.update({
                "date": day.isoformat(),
                "description": description[:255],
                "amount": format(abs(signed), "f"),
                "direction": direction,
                "type": guess_type(description, direction),
                "merchant": entities.get("merchant"),
                "payment_method": payment_method(entities),
                "entities": entities,
                "external_id": col(row, "external_id").strip() or None,
                "notes": col(row, "notes").strip(),
                "balance": col(row, "balance").strip(),
                "status": "ok",
            })
        except (ValueError, TypeError) as exc:
            item.update({"status": "error", "error": str(exc)})
        out.append(item)
    return out
