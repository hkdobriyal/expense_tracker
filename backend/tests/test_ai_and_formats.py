"""Entity extraction, every import format (incl. OCR), the ML categoriser and the assistant."""

import io
import json
import threading
from datetime import date
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

from app.services.extraction import extract
from app.services.statement_text import parse_lines

from .conftest import add_txn, category_id, make_account

# ---------------------------------------------------------------------------
# Entity extraction
# ---------------------------------------------------------------------------

EXTRACTION_CASES = [
    ("UPI/DR/425167812345/SWIGGY LIMITED/YESB/swiggy8@ybl/Food order", {"mode": "UPI", "merchant": "Swiggy", "reference": "425167812345", "vpa": "swiggy8@ybl", "app": "PhonePe", "direction": "out"}),
    ("UPI/426512345678/Payment from Ph/9876543210@ybl/HDFC BANK", {"mode": "UPI", "counterparty": "person", "merchant": None}),
    ("UPI/P2A/426512349999/RAJESH KUMAR/SBIN/rajesh.k@oksbi/rent sept", {"payee": "Rajesh Kumar", "counterparty": "person", "note": "rent sept", "app": "Google Pay"}),
    ("UPI-ZOMATO LTD-ZOMATO@PTYBL-YESB0PTMUPI-425198765432-PAYMENT FROM PHONE", {"mode": "UPI", "merchant": "Zomato", "bank": "Yes Bank", "reference": "425198765432"}),
    ("NEFT CR-ICIC0000104-ACME TECHNOLOGIES PVT LTD-SALARY SEP 2026-ICICN52026091500123", {"mode": "NEFT", "direction": "in", "merchant": "Acme", "bank": "ICICI Bank"}),
    ("IMPS/P2A/426598761234/PRIYA SHARMA/HDFC/XX1234/Trip share", {"mode": "IMPS", "payee": "Priya Sharma", "reference": "426598761234"}),
    ("POS 438976XXXXXX1234 DECATHLON SPORTS BANGALORE", {"mode": "Card", "merchant": "Decathlon", "card_last4": "1234", "city": "Bangalore"}),
    ("ACH D- ZERODHA BROKING LTD-ZERODHA123456", {"mode": "Auto-debit", "merchant": "Zerodha"}),
    ("RAZORPAY*URBANCLAP TECHNOLOGIES", {"merchant": "Urbanclap"}),
    ("INT.PD:1234567890:01-07-2026 TO 30-09-2026", {"mode": "Interest", "direction": "in"}),
]


@pytest.mark.parametrize("text,expected", EXTRACTION_CASES)
def test_entity_extraction(text, expected):
    got = extract(text)
    for key, value in expected.items():
        assert got.get(key) == value, (key, got)


# ---------------------------------------------------------------------------
# Free-text statements (PDF text / DOCX / OCR share this parser)
# ---------------------------------------------------------------------------

STATEMENT_LINES = """
STATE BANK OF INDIA  Statement of Account
Opening Balance 12,500.00
Date        Value Date   Description                                  Debit      Credit     Balance
01-09-2026  01-09-2026   NEFT CR-ICIC0000104-ACME TECH-SALARY          85,000.00  97,500.00
03/09/2026  03/09/2026   UPI/DR/425167812345/SWIGGY LIMITED/YESB/      340.00     97,160.00
                         swiggy8@ybl/Food order
05 Sep 2026              ACH D- ZERODHA BROKING LTD                    10,000.00  87,160.00
07-Sep-26                UPI/429999999999/AMAZON REFUND/REVERSAL      499.00     87,659.00
Page 1 of 1
Closing Balance 87,659.00
""".strip().splitlines()


def test_free_text_statement_infers_direction_from_balance():
    rows, notes = parse_lines(STATEMENT_LINES)
    assert [(r[0][:6], r[2], r[3]) for r in rows] == [("01-09-", "85000.00", "CR"), ("03/09/", "340.00", "DR"), ("05 Sep", "10000.00", "DR"), ("07-Sep", "499.00", "CR")]
    assert "swiggy8@ybl/Food order" in rows[1][1]  # continuation line joined
    assert "running balance" in notes[0]


def upload(api, filename, content, account_id, password=None):
    data = {"account_id": str(account_id)}
    if password:
        data["password"] = password
    return api.post("/api/imports", files={"file": (filename, content, "application/octet-stream")}, data=data)


def _minimal_pdf(lines: list[str]) -> bytes:
    """Tiny hand-built PDF with real text (no PDF library needed to create it)."""
    y, ops = 760, []
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"BT /F1 9 Tf 30 {y} Td ({safe}) Tj ET")
        y -= 14
    stream = "\n".join(ops).encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offsets))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_text_pdf_and_password_protected_pdf(api):
    acc = make_account(api, opening="12500")
    pdf = _minimal_pdf(STATEMENT_LINES)
    job = upload(api, "sbi.pdf", pdf, acc["id"])
    assert job["parse_method"] == "pdf-text" and job["counts"]["ok"] == 4 and job["bank_detected"] == "State Bank of India"
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(pdf)).pages:
        writer.add_page(page)
    writer.encrypt("ABCDE1234F")
    buf = io.BytesIO()
    writer.write(buf)
    api.post("/api/imports", files={"file": ("locked.pdf", buf.getvalue(), "application/pdf")}, data={"account_id": str(acc["id"])}, expected=422)
    job = upload(api, "locked.pdf", buf.getvalue(), acc["id"], password="ABCDE1234F")
    assert job["counts"]["duplicates"] == 0 and job["counts"]["ok"] == 4
    result = api.post(f"/api/imports/{job['id']}/commit", json={})
    assert result["imported"] == 4
    txns = {t["description"]: t for t in api.get("/api/transactions")["items"]}
    swiggy = txns["Swiggy"]
    assert swiggy["extracted"]["vpa"] == "swiggy8@ybl" and swiggy["payment_method"] == "UPI" and swiggy["category"] == "Food delivery"
    assert next(a for a in api.get("/api/accounts") if a["id"] == acc["id"])["balance_minor"] == 8_765_900


def test_docx_html_json_and_txt_imports(api):
    import docx

    acc = make_account(api)
    d = docx.Document()
    d.add_paragraph("HDFC BANK statement")
    table = d.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["Date", "Narration", "Withdrawal Amt", "Deposit Amt"]):
        cell.text = text
    for row in [("02/09/2026", "UPI/DR/111122223333/ZEPTO/YESB/zepto@ybl/groceries", "450.00", ""), ("03/09/2026", "NEFT CR-HDFC0000001-CLIENT CO-INVOICE 12", "", "25,000.00")]:
        cells = table.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    buf = io.BytesIO()
    d.save(buf)
    job = upload(api, "statement.docx", buf.getvalue(), acc["id"])
    assert job["file_format"] == "docx" and job["parse_method"] == "docx-table" and job["counts"]["ok"] == 2

    html = b"<html><body><table><tr><th>Txn Date</th><th>Description</th><th>Debit</th><th>Credit</th></tr><tr><td>04/09/2026</td><td>POS 4389XXXXXXXX1234 CROMA BLR</td><td>12,999.00</td><td></td></tr></table></body></html>"
    job = upload(api, "statement.xls", html, acc["id"])  # banks' .xls downloads are often HTML
    assert job["file_format"] == "xls" and job["counts"]["ok"] == 1 and job["rows"][0]["merchant"] == "Croma"

    payload = json.dumps([{"date": "2026-09-05", "description": "Netflix", "amount": "-649.00"}]).encode()
    job = upload(api, "export.json", payload, acc["id"])
    assert job["counts"]["ok"] == 1 and job["rows"][0]["direction"] == "out"

    job = upload(api, "notes.txt", "\n".join(STATEMENT_LINES).encode(), acc["id"])
    assert job["parse_method"] == "text" and job["counts"]["ok"] == 4


def test_scanned_statement_image_via_ocr(api):
    from PIL import Image, ImageDraw, ImageFont

    acc = make_account(api, opening="5000")
    try:  # a real font, like a scanned statement (Pillow's tiny bitmap font is not representative)
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        try:
            font = ImageFont.truetype("DejaVuSans.ttf", 34)
        except OSError:
            pytest.skip("no TrueType font available to render a test scan")
    img = Image.new("RGB", (1500, 300), "white")
    draw = ImageDraw.Draw(img)
    lines = ["Opening Balance 5,000.00", "10/09/2026    UPI BLINKIT GROCERIES     640.00     4,360.00", "12/09/2026    SALARY CREDIT ACME     50,000.00     54,360.00"]
    for i, line in enumerate(lines):
        draw.text((30, 30 + i * 80), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    job = upload(api, "scan.png", buf.getvalue(), acc["id"])
    assert job["parse_method"] == "image-ocr"
    rows = [r for r in job["rows"] if r["status"] == "ok"]
    assert [(r["amount"], r["direction"]) for r in rows] == [("640.00", "out"), ("50000.00", "in")]


# ---------------------------------------------------------------------------
# ML categoriser
# ---------------------------------------------------------------------------

def test_ml_learns_your_own_merchants(api):
    acc = make_account(api)
    coffee = category_id(api, "Coffee & snacks")
    for i in range(6):
        add_txn(api, acc["id"], str(80 + i), description=f"UPI/DR/42{i}000000001/CHAI SUTTA BAR/YESB/chaisutta{i}@ybl/chai", category_id=coffee, day=date(2026, 9, i + 1))
    api.post("/api/ai/train")
    txn = add_txn(api, acc["id"], "95", description="UPI/DR/429999999999/CHAI SUTTA BAR KORAMANGALA/YESB/chaisuttablr@ybl/tea")
    assert txn["category_id"] == coffee and txn["category_source"] in ("ml", "merchant")
    suggestion = api.post("/api/ai/suggest", json={"description": "chai sutta bar indiranagar", "amount": "60"})
    assert suggestion["suggestions"][0]["category_id"] == coffee
    status = api.get("/api/ai/status")
    assert status["ml"]["trained"] is True and status["llm"]["enabled"] is False


def test_ml_works_from_day_one_with_seed_examples(api):
    acc = make_account(api)
    txn = add_txn(api, acc["id"], "250", description="UPI/DR/425100000000/UBER INDIA/YESB/uber.rides@axl/trip")
    assert txn["category"] == "Cab & auto"
    assert txn["category_source"] in ("ml", "keyword")


# ---------------------------------------------------------------------------
# Assistant – rules engine and a fake local LLM
# ---------------------------------------------------------------------------

def _ledger(api):
    acc = make_account(api, opening="100000")
    add_txn(api, acc["id"], "90000", "income", "Salary", category_id=category_id(api, "Salary"), day=date(2026, 9, 1))
    add_txn(api, acc["id"], "1200", description="Dinner", category_id=category_id(api, "Restaurants"), day=date(2026, 9, 10))
    add_txn(api, acc["id"], "800", description="Lunch", category_id=category_id(api, "Restaurants"), day=date(2026, 8, 10))
    add_txn(api, acc["id"], "3000", description="Groceries", category_id=category_id(api, "Groceries"), day=date(2026, 8, 12))
    api.post("/api/subscriptions", json={"name": "Netflix", "amount": "649", "frequency": "monthly"})
    return acc


def test_assistant_rules_engine(api):
    _ledger(api)
    ask = lambda q: api.post("/api/ai/ask", json={"question": q})  # noqa: E731
    r = ask("How much did I spend on restaurants in September 2026?")
    assert r["engine"] == "rules" and "₹1,200" in r["answer"] and "Restaurants" in r["answer"]
    r = ask("Compare August 2026 and September 2026")
    assert "₹3,800 in August 2026" in r["answer"] and "₹1,200 in September 2026" in r["answer"]
    assert "₹649" in ask("Which subscriptions cost the most?")["answer"]
    assert "₹90,000" in ask("How much did I earn in September 2026?")["answer"]


class _FakeLLM(BaseHTTPRequestHandler):
    """Speaks just enough of the OpenAI chat API: first a tool call, then a final answer."""

    def log_message(self, *args):
        pass

    def _send(self, payload):
        body = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        self._send({"data": [{"id": "qwen2.5:3b"}]})

    def do_POST(self):
        req = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if req.get("response_format"):  # categorisation request
            items = json.loads(req["messages"][1]["content"])["items"]
            self._send({"choices": [{"message": {"content": json.dumps({"results": [{"id": i["id"], "category": "Health/Pharmacy", "merchant": "Local Chemist", "confidence": 0.8} for i in items]})}}]})
        elif not any(m["role"] == "tool" for m in req["messages"]):
            self._send({"choices": [{"message": {"content": "", "tool_calls": [{"id": "c1", "type": "function", "function": {"name": "spending", "arguments": json.dumps({"period": "2026-09", "category": "Restaurants"})}}]}}]})
        else:
            tool = json.loads(next(m for m in req["messages"] if m["role"] == "tool")["content"])
            self._send({"choices": [{"message": {"content": f"You spent {tool['spent_text']} on restaurants in {tool['period']}."}}]})


@pytest.fixture
def fake_llm(monkeypatch):
    from app.config import get_settings
    from app.services import llm

    server = HTTPServer(("127.0.0.1", 0), _FakeLLM)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    monkeypatch.setenv("LLM_ENABLED", "true")
    monkeypatch.setenv("LLM_BASE_URL", f"http://127.0.0.1:{server.server_port}/v1")
    monkeypatch.setenv("LLM_MODEL", "qwen2.5:3b")
    get_settings.cache_clear()
    llm._status_cache.update(at=0.0, value=None)
    yield
    server.shutdown()
    monkeypatch.delenv("LLM_ENABLED")
    get_settings.cache_clear()
    llm._status_cache.update(at=0.0, value=None)


def test_assistant_with_local_llm_uses_ledger_numbers(api, fake_llm):
    _ledger(api)
    r = api.post("/api/ai/ask", json={"question": "what did restaurants cost me in september?"})
    assert r["engine"] == "llm" and r["tools_used"] == ["spending"]
    assert r["answer"] == "You spent ₹1,200 on restaurants in September 2026."


def test_llm_categorizes_what_ml_is_unsure_about(api, fake_llm, monkeypatch):
    from app.services import ml

    acc = make_account(api)
    txn = add_txn(api, acc["id"], "180", description="QWERTY ZXCV 998877")
    monkeypatch.setattr(ml, "predict", lambda *a, **k: [])  # force "ML unsure"
    api.patch(f"/api/transactions/{txn['id']}", json={"category_id": None})
    result = api.post("/api/ai/categorize", json={"scope": "uncategorized"})
    assert result["categorized_by_llm"] == 1
    t = api.get(f"/api/transactions/{txn['id']}")
    assert (t["category"], t["category_source"]) == ("Pharmacy", "llm")


def test_migration_upgrades_a_database_with_existing_rows(tmp_path):
    import sqlite3

    from alembic import command

    from app.migrations import alembic_config

    url = f"sqlite:///{(tmp_path / 'old.sqlite3').as_posix()}"
    cfg = alembic_config(url)
    command.upgrade(cfg, "88ba7aefc32f")  # the first release
    con = sqlite3.connect(tmp_path / "old.sqlite3")
    con.execute("insert into users (id, email, password_hash, display_name, is_demo, created_at) values (1,'a@b.co','x','A',0,'2026-01-01')")
    con.execute("insert into accounts (id,user_id,name,type,institution,currency,opening_balance_minor,account_number_mask,color,include_in_net_worth,is_archived,created_at,updated_at) values (1,1,'Bank','savings','','INR',0,'','',1,0,'2026-01-01','2026-01-01')")
    con.execute("insert into transactions (user_id,account_id,type,date,amount_minor,currency,base_amount_minor,description,notes,payment_method,is_recurring,reviewed,source,fingerprint,raw_description,created_at,updated_at) values (1,1,'expense','2026-01-02',100,'INR',100,'Tea','','',0,1,'manual','fp','', '2026-01-02','2026-01-02')")
    con.commit()
    con.close()
    command.upgrade(cfg, "head")
    con = sqlite3.connect(tmp_path / "old.sqlite3")
    assert con.execute("select extracted, category_source from transactions").fetchone() == ("{}", "")
